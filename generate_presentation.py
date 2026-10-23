"""
Think Before You Click: Phishing, Scams & Social Engineering
PowerPoint generator
Presenter: Nikhil Jamkar

Install dependencies:
    pip install -r requirements.txt

Run:
    python generate_presentation.py

Output:
    Think_Before_You_Click_Nikhil_Jamkar.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUTPUT_FILE = "Think_Before_You_Click_Nikhil_Jamkar.pptx"

NAVY = RGBColor(15, 29, 55)
BLUE = RGBColor(39, 105, 190)
CYAN = RGBColor(77, 210, 218)
WHITE = RGBColor(247, 250, 255)
MUTED = RGBColor(190, 204, 225)
ORANGE = RGBColor(255, 177, 66)
GREEN = RGBColor(74, 201, 145)


def set_bg(slide, color=NAVY):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_round_rect(slide, x, y, w, h, fill_color, line_color=None, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = line_color or fill_color
    return shape


def add_header(slide, number, title, subtitle=None):
    add_text(slide, f"{number:02d}  /  CYBER SMART", 0.65, 0.38, 4.0, 0.3,
             11, CYAN, True)
    add_text(slide, title, 0.65, 0.82, 12.0, 0.62, 30, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.48, 11.7, 0.45, 14, MUTED)
    add_round_rect(slide, 0.68, 2.03, 0.9, 0.06, CYAN, CYAN, False)


def add_bullet(slide, text, x, y, w=7.5, size=19, color=WHITE,
               bullet_color=CYAN):
    add_round_rect(slide, x, y + 0.11, 0.12, 0.12, bullet_color, bullet_color)
    add_text(slide, text, x + 0.28, y, w - 0.28, 0.56, size, color)


def add_footer(slide, number):
    add_text(slide, "THINK BEFORE YOU CLICK  •  STAY ALERT, STAY SAFE",
             0.68, 7.12, 9.8, 0.2, 9, MUTED)
    add_text(slide, str(number), 12.15, 7.08, 0.45, 0.25,
             10, CYAN, True, align=PP_ALIGN.RIGHT)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: Title
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    for x, y, diameter, color in [
        (9.7, 1.15, 2.25, BLUE),
        (10.75, 2.35, 1.55, CYAN),
        (9.25, 3.55, 1.0, ORANGE),
    ]:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y),
            Inches(diameter), Inches(diameter)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.fill.transparency = 18
        shape.line.fill.background()
    add_text(slide, "CYBERSECURITY AWARENESS", 0.85, 0.8, 5.8, 0.35,
             13, CYAN, True)
    add_text(slide, "THINK BEFORE\nYOU CLICK", 0.8, 1.45, 8.4, 2.05,
             40, WHITE, True)
    add_text(slide, "Phishing • Scams • Social Engineering",
             0.85, 3.75, 7.5, 0.55, 22, MUTED)
    add_round_rect(slide, 0.85, 4.7, 4.25, 0.78, BLUE)
    add_text(slide, "Presented by  Nikhil Jamkar",
             1.08, 4.91, 3.85, 0.3, 17, WHITE, True)
    add_text(slide, "One careful moment can protect your identity, money and data.",
             0.85, 6.25, 8.2, 0.42, 15, MUTED)
    add_text(slide, "01", 12.1, 7.05, 0.5, 0.25, 10, CYAN, True,
             align=PP_ALIGN.RIGHT)

    # Slide 2: Why it matters
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 2, "Why this topic matters",
               "Every click is a choice—make it a safe one.")
    add_bullet(slide, "Our daily lives depend on email, messaging, apps and online payments.",
               0.85, 2.45, 8.1)
    add_bullet(slide, "Cybercriminals use fake messages and clever persuasion to gain access.",
               0.85, 3.35, 8.1)
    add_bullet(slide, "A single rushed action can expose passwords, money or personal information.",
               0.85, 4.25, 8.1)
    add_round_rect(slide, 9.45, 2.55, 2.65, 2.35, BLUE)
    add_text(slide, "PAUSE", 9.7, 2.92, 2.15, 0.55, 28, WHITE, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "CHECK", 9.7, 3.67, 2.15, 0.45, 20, CYAN, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "THEN ACT", 9.7, 4.25, 2.15, 0.45, 20, WHITE, True,
             align=PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # Slide 3: Phishing
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 3, "What is phishing?",
               "A fake communication designed to trick you into taking an unsafe action.")
    add_bullet(slide, "It may arrive by email, SMS, social media, phone call or messaging app.",
               0.85, 2.4, 8.1)
    add_bullet(slide, "The sender may pretend to be a bank, delivery service, college or trusted person.",
               0.85, 3.25, 8.1)
    add_bullet(slide, "The goal may be to steal login details, payment information or install harmful software.",
               0.85, 4.1, 8.1)
    add_round_rect(slide, 9.45, 2.55, 2.55, 2.25, RGBColor(31, 52, 83))
    add_text(slide, "FAKE", 9.8, 2.88, 1.85, 0.5, 28, ORANGE, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "MESSAGE", 9.8, 3.55, 1.85, 0.5, 21, WHITE, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "→", 10.25, 4.15, 1.0, 0.45, 26, CYAN, True,
             align=PP_ALIGN.CENTER)
    add_footer(slide, 3)

    # Slide 4: Scams
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 4, "Common online scams",
               "Recognising the pattern helps you stay in control.")
    cards = [
        ("BANK / KYC", "A message says your account will be blocked unless you verify details.", ORANGE),
        ("PRIZE / REWARD", "You are told you won something—but must pay a fee first.", CYAN),
        ("JOB / INVESTMENT", "A promise of easy money pressures you to pay or share details.", GREEN),
    ]
    for index, (title, description, color) in enumerate(cards):
        x = 0.8 + index * 4.15
        add_round_rect(slide, x, 2.55, 3.7, 2.65, RGBColor(27, 45, 74))
        add_round_rect(slide, x + 0.25, 2.85, 0.55, 0.08, color, color, False)
        add_text(slide, title, x + 0.25, 3.18, 3.2, 0.45, 17, color, True)
        add_text(slide, description, x + 0.25, 3.82, 3.15, 1.05, 16, WHITE)
    add_text(slide, "If it sounds urgent, too good to be true, or asks for secrecy—pause and verify.",
             0.85, 5.75, 11.5, 0.6, 18, WHITE, True)
    add_footer(slide, 4)

    # Slide 5: Social engineering
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 5, "Social engineering: the human angle",
               "Attackers often target trust and emotions—not just technology.")
    items = [
        ("URGENCY", "“Act now or your account will close.”"),
        ("AUTHORITY", "“I’m from support—share the OTP.”"),
        ("FEAR", "“There is a problem with your account.”"),
        ("CURIOSITY / REWARD", "“Open this surprise file to claim a gift.”"),
    ]
    for index, (label, example) in enumerate(items):
        y = 2.35 + index * 0.88
        add_round_rect(slide, 0.9, y, 2.55, 0.58, BLUE)
        add_text(slide, label, 1.05, y + 0.12, 2.25, 0.25, 14, WHITE, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, example, 3.8, y + 0.09, 7.8, 0.38, 18, WHITE)
    add_text(slide, "A genuine organisation will not ask you to reveal your password or OTP.",
             0.9, 6.2, 11.4, 0.4, 17, CYAN, True)
    add_footer(slide, 5)

    # Slide 6: Warning signs
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 6, "Warning signs to look for",
               "Small clues can reveal a suspicious message.")
    left = [
        "Unexpected links or attachments",
        "Spelling mistakes or unusual sender address",
        "Requests for passwords, OTPs or payment",
    ]
    right = [
        "Threats, deadlines or pressure to act fast",
        "Web addresses that imitate real brands",
        "Offers that seem unrealistically generous",
    ]
    for index, item in enumerate(left):
        add_bullet(slide, item, 0.9, 2.45 + index * 0.95, 5.65, 17)
    for index, item in enumerate(right):
        add_bullet(slide, item, 6.9, 2.45 + index * 0.95, 5.5, 17, WHITE, ORANGE)
    add_round_rect(slide, 2.1, 5.75, 9.0, 0.62, RGBColor(31, 52, 83))
    add_text(slide, "Check the sender and website independently—never rely on the message alone.",
             2.35, 5.93, 8.55, 0.27, 15, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)

    # Slide 7: Safe habits
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 7, "Think before you click: 5 safe habits",
               "Build a quick safety check into your routine.")
    habits = [
        ("1", "PAUSE", "Do not act under pressure."),
        ("2", "INSPECT", "Check the sender and full web address."),
        ("3", "VERIFY", "Contact the organisation using its official app or number."),
        ("4", "PROTECT", "Use strong, unique passwords and multi-factor authentication."),
        ("5", "REPORT", "Flag suspicious messages and warn others."),
    ]
    for index, (number, title, description) in enumerate(habits):
        y = 2.25 + index * 0.82
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.95), Inches(y), Inches(0.48), Inches(0.48)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = BLUE
        shape.line.fill.background()
        add_text(slide, number, 1.05, y + 0.1, 0.28, 0.22, 13, WHITE, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, title, 1.75, y + 0.04, 1.65, 0.3, 16, CYAN, True)
        add_text(slide, description, 3.55, y + 0.03, 8.5, 0.36, 17, WHITE)
    add_footer(slide, 7)

    # Slide 8: Account and device protection
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 8, "Protect your accounts and devices",
               "Prevention is easier when good security is switched on.")
    bullets = [
        "Use a different strong password for each important account; consider a password manager.",
        "Enable multi-factor authentication (MFA), especially for email and banking.",
        "Keep your phone, computer, browser and apps updated.",
        "Download apps only from trusted sources and review permissions.",
        "Never share OTPs, PINs or passwords—even with someone claiming to be support.",
    ]
    for index, item in enumerate(bullets):
        add_bullet(slide, item, 0.9, 2.35 + index * 0.83, 11.2, 17,
                   WHITE, ORANGE if index == 4 else CYAN)
    add_footer(slide, 8)

    # Slide 9: Response
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_header(slide, 9, "What to do if you clicked or shared details",
               "Act quickly. A mistake can still be contained.")
    steps = [
        ("1", "Disconnect", "If a file seems harmful, disconnect the device from the network."),
        ("2", "Secure", "Change affected passwords from a trusted device; enable MFA."),
        ("3", "Contact", "Notify your bank or service provider immediately if money or account access is at risk."),
        ("4", "Report", "Tell your IT team or platform, save evidence, and report the incident through official channels."),
    ]
    for index, (number, title, description) in enumerate(steps):
        x = 0.85 + (index % 2) * 6.15
        y = 2.45 + (index // 2) * 1.7
        add_round_rect(slide, x, y, 5.55, 1.35, RGBColor(27, 45, 74))
        add_text(slide, number, x + 0.2, y + 0.18, 0.42, 0.4, 22, CYAN, True)
        add_text(slide, title, x + 0.8, y + 0.18, 4.3, 0.34, 18, WHITE, True)
        add_text(slide, description, x + 0.8, y + 0.6, 4.35, 0.58, 13, MUTED)
    add_text(slide, "In India, report cybercrime through the official National Cyber Crime Reporting Portal; call 1930 promptly for financial cyber fraud.",
             0.9, 6.0, 11.7, 0.62, 13, ORANGE, True)
    add_footer(slide, 9)

    # Slide 10: Summary
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, "10  /  KEY TAKEAWAY", 0.7, 0.45, 4.2, 0.3, 11, CYAN, True)
    add_text(slide, "Your best defence\nis a thoughtful pause.",
             0.8, 1.1, 8.9, 1.65, 35, WHITE, True)
    add_text(slide, "Before you click, ask:", 0.85, 3.2, 5.2, 0.45, 21, MUTED, True)
    questions = [
        "Do I trust the sender?",
        "Is this link or request expected?",
        "Can I verify it another way?",
    ]
    for index, question in enumerate(questions):
        y = 3.9 + index * 0.58
        add_round_rect(slide, 0.9, y, 0.18, 0.18, CYAN)
        add_text(slide, question, 1.25, y - 0.04, 6.5, 0.3, 18, WHITE)
    add_round_rect(slide, 8.8, 2.25, 3.25, 2.4, BLUE)
    add_text(slide, "STOP  •  CHECK\nTHEN CLICK", 9.05, 2.9, 2.75, 1.1,
             22, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Thank you!", 0.85, 6.15, 3.0, 0.55, 24, CYAN, True)
    add_text(slide, "Presented by Nikhil Jamkar", 0.85, 6.68, 4.8, 0.3, 14, MUTED)
    add_text(slide, "10", 12.15, 7.05, 0.45, 0.25, 10, CYAN, True,
             align=PP_ALIGN.RIGHT)

    prs.save(OUTPUT_FILE)
    print(f"Presentation created: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
